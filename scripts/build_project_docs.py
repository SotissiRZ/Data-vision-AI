#!/usr/bin/env python3
from __future__ import annotations
import subprocess
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables' / 'documentation'
OUT.mkdir(parents=True, exist_ok=True)
VERSION = (ROOT / 'VERSION').read_text(encoding='utf-8').strip()
DOCS = [
    ('docs/RAPPORT_TECHNIQUE.md', f'Rapport_Technique_DataVision_v{VERSION}.docx', 'Rapport technique'),
    ('docs/GUIDE_UTILISATEUR.md', f'Guide_Utilisateur_DataVision_v{VERSION}.docx', 'Guide utilisateur'),
    ('docs/GUIDE_DEPLOIEMENT.md', f'Guide_Deploiement_DataVision_v{VERSION}.docx', 'Guide de déploiement'),
    ('docs/RAPPORT_SECURITE.md', f'Rapport_Securite_DataVision_v{VERSION}.docx', 'Rapport de sécurité'),
    ('docs/GUIDE_EXPLOITATION.md', f'Guide_Exploitation_DataVision_v{VERSION}.docx', "Guide d'exploitation"),
]

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def postprocess(path: Path, short_title: str):
    doc = Document(path)
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    for section in doc.sections:
        section.header_distance = Cm(0.8); section.footer_distance = Cm(0.8)
        hp = section.header.paragraphs[0]
        hp.text = f'DataVision AI v{VERSION} — {short_title}'
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if hp.runs:
            hp.runs[0].font.size = Pt(8); hp.runs[0].font.name = 'Aptos'
        fp = section.footer.paragraphs[0]
        add_page_number(fp)
        for r in fp.runs:
            r.font.size = Pt(8); r.font.name = 'Aptos'
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
    for name, size in [('Title',24),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
        if name in styles:
            styles[name].font.name = 'Aptos'; styles[name].font.size = Pt(size)
    first_content = next((p for p in doc.paragraphs if p.text.strip()), None)
    if first_content is not None:
        first_content.style = styles['Title']
    for p in doc.paragraphs:
        if p.style.name == 'Title':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
        if p.style.name.startswith('Heading'):
            p.paragraph_format.keep_with_next = True
    for table in doc.tables:
        if 'Table Grid' in [st.name for st in doc.styles if getattr(st, 'type', None) == 3]:
            table.style = 'Table Grid'
        if table.rows:
            for cell in table.rows[0].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True
    cp = doc.core_properties
    cp.author = 'DataVision AI Project'
    cp.title = short_title
    cp.subject = f'DataVision AI v{VERSION}'
    cp.comments = 'Document généré pour la release DataVision AI.'
    doc.save(path)

for src, filename, title in DOCS:
    srcp = ROOT / src
    outp = OUT / filename
    subprocess.run(['pandoc', str(srcp), '-o', str(outp), '--from=gfm'], check=True)
    postprocess(outp, title)

index = OUT / 'README.txt'
index.write_text(
    f'DataVision AI v{VERSION} - Dossier documentaire\n\n'
    'Livrables : rapport technique, guide utilisateur, guide de déploiement, rapport de sécurité, guide d’exploitation.\n'
    'Les versions PDF sont générées après validation visuelle des DOCX.\n',
    encoding='utf-8'
)
print(OUT)
