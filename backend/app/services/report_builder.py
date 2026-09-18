from __future__ import annotations

import html
import json
import math
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.graphics import renderPM, renderSVG
from reportlab.graphics.shapes import Circle, Drawing, Line, Polygon, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.core.config import get_settings
from app.services.analysis_history import get_analysis
from app.services.decision import decision_support
from app.services.profiling import profile_dataframe
from app.services.quality import quality_report
from app.services.saved_visualizations import list_visualizations
from app.services.visualization import build_visualization
from app.services.storage import get_meta, load_dataframe

BRAND = {
    "navy": "#203C4B",
    "teal": "#2798B7",
    "teal_dark": "#147A98",
    "aqua": "#E7F6FA",
    "ink": "#243844",
    "muted": "#687E89",
    "line": "#D8E4E9",
    "surface": "#F5F8FA",
    "success": "#178A63",
    "warning": "#B77A14",
    "danger": "#B44A53",
}

TEMPLATES = {
    "executive": {
        "label": "Executif",
        "description": "Synthese decisionnelle, resultats prioritaires, graphiques essentiels et recommandations.",
        "default_sections": ["executive_summary", "analytical_story", "visualizations", "limitations", "methodology", "provenance"],
    },
    "analytical": {
        "label": "Analytique",
        "description": "Rapport complet avec narration analytique, statistiques, graphiques et interpretation.",
        "default_sections": ["executive_summary", "analytical_story", "overview", "quality", "descriptive", "visualizations", "ai_analysis", "limitations", "methodology", "provenance"],
    },
    "technical": {
        "label": "Technique",
        "description": "Version detaillee orientee audit, methodologie, limites et reproductibilite.",
        "default_sections": ["overview", "quality", "descriptive", "visualizations", "ai_analysis", "analytical_story", "limitations", "methodology", "provenance"],
    },
}

SECTION_LABELS = {
    "executive_summary": "Synthese executive",
    "analytical_story": "Resultats cles et lecture analytique",
    "overview": "Vue d'ensemble des donnees",
    "quality": "Qualite des donnees",
    "descriptive": "Statistiques descriptives",
    "visualizations": "Analyses visuelles",
    "ai_analysis": "Analyse AI Analyst",
    "limitations": "Limites et precautions d'interpretation",
    "methodology": "Methodologie",
    "provenance": "Provenance et reproductibilite",
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _dir() -> Path:
    path = get_settings().data_root / "reports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _slug(value: str) -> str:
    safe = re.sub(r"[^a-zA-Z0-9_-]+", "-", value.strip()).strip("-").lower()
    return safe[:80] or "rapport"


def _numeric_summary(df: pd.DataFrame, max_columns: int = 16) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for column in [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c])][:max_columns]:
        s = pd.to_numeric(df[column], errors="coerce").dropna()
        if s.empty:
            continue
        rows.append({
            "variable": str(column),
            "n": int(s.size),
            "mean": float(s.mean()),
            "median": float(s.median()),
            "std": float(s.std(ddof=1)) if s.size > 1 else 0.0,
            "q1": float(s.quantile(.25)),
            "q3": float(s.quantile(.75)),
            "min": float(s.min()),
            "max": float(s.max()),
        })
    return rows


def _type_summary(profile: dict[str, Any]) -> list[dict[str, Any]]:
    buckets = {"Numerique": 0, "Texte / categorie": 0, "Date / temps": 0, "Autre": 0}
    for col in profile.get("columns", []):
        dtype = str(col.get("dtype", "")).lower()
        if any(x in dtype for x in ["int", "float", "double", "decimal"]):
            buckets["Numerique"] += 1
        elif any(x in dtype for x in ["date", "time"]):
            buckets["Date / temps"] += 1
        elif any(x in dtype for x in ["object", "string", "category", "bool"]):
            buckets["Texte / categorie"] += 1
        else:
            buckets["Autre"] += 1
    return [{"type": k, "count": v} for k, v in buckets.items() if v]


def _missing_summary(profile: dict[str, Any], limit: int = 12) -> list[dict[str, Any]]:
    rows = [
        {"variable": c.get("name"), "missing": int(c.get("missing", 0)), "missing_pct": float(c.get("missing_pct", 0.0))}
        for c in profile.get("columns", []) if int(c.get("missing", 0)) > 0
    ]
    rows.sort(key=lambda x: x["missing_pct"], reverse=True)
    return rows[:limit]


def _severity_counts(quality: dict[str, Any]) -> dict[str, int]:
    out = {"high": 0, "medium": 0, "low": 0}
    for issue in quality.get("issues", []):
        sev = str(issue.get("severity", "low")).lower()
        out[sev if sev in out else "low"] += 1
    return out


def _fmt(value: Any) -> str:
    if isinstance(value, float):
        if math.isnan(value):
            return "-"
        if abs(value) >= 1000:
            return f"{value:,.2f}".replace(",", " ")
        return f"{value:.4g}"
    if value is None:
        return "-"
    return str(value)


def _format_date(value: str | None) -> str:
    if not value:
        return "-"
    try:
        d = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return d.strftime("%d/%m/%Y %H:%M")
    except Exception:
        return str(value)


def _executive_summary(profile: dict[str, Any], quality: dict[str, Any], decision: dict[str, Any], analysis: dict[str, Any] | None) -> dict[str, Any]:
    total_cells = max(int(profile.get("rows", 0)) * int(profile.get("columns_count", 0)), 1)
    missing_cells = sum(int(c.get("missing", 0)) for c in profile.get("columns", []))
    missing_pct = 100.0 * missing_cells / total_cells
    severity = _severity_counts(quality)
    highlights: list[dict[str, str]] = []
    if quality.get("score", 0) >= 90:
        highlights.append({"level": "positive", "title": "Qualite globale elevee", "text": f"Score de qualite {quality.get('score')}/100."})
    elif quality.get("score", 0) >= 70:
        highlights.append({"level": "attention", "title": "Qualite acceptable avec corrections", "text": f"Score de qualite {quality.get('score')}/100 et {quality.get('issues_count', 0)} alerte(s)."})
    else:
        highlights.append({"level": "critical", "title": "Qualite a traiter avant modelisation", "text": f"Score de qualite {quality.get('score')}/100."})
    if missing_pct > 0:
        highlights.append({"level": "attention" if missing_pct < 10 else "critical", "title": "Valeurs manquantes", "text": f"{missing_cells} cellule(s) manquante(s), soit {missing_pct:.1f}% du volume."})
    if profile.get("duplicates"):
        highlights.append({"level": "attention", "title": "Doublons detectes", "text": f"{profile.get('duplicates')} ligne(s) dupliquee(s) a verifier."})
    if analysis:
        for finding in analysis.get("findings", [])[:3]:
            highlights.append({"level": str(finding.get("level", "info")), "title": str(finding.get("title", "Constat analytique")), "text": str(finding.get("statement", ""))})
    recommendations = [
        {"priority": a.get("priority", "medium"), "action": a.get("action", ""), "evidence": a.get("evidence", "")}
        for a in decision.get("actions", [])[:5]
    ]
    return {
        "kpis": [
            {"label": "Observations", "value": int(profile.get("rows", 0)), "hint": "lignes"},
            {"label": "Variables", "value": int(profile.get("columns_count", 0)), "hint": "colonnes"},
            {"label": "Qualite", "value": int(quality.get("score", 0)), "hint": "/100"},
            {"label": "Manquants", "value": round(missing_pct, 1), "hint": "% des cellules"},
        ],
        "highlights": highlights[:6],
        "recommendations": recommendations,
        "severity": severity,
    }




def _is_identifier_series(name: str, series: pd.Series, rows: int) -> bool:
    n = str(name).lower()
    named = n in {"id", "index", "row", "record", "patient_id", "customer_id", "user_id"} or n.endswith("_id")
    unique = int(series.nunique(dropna=True))
    high_unique = rows > 20 and unique / max(rows, 1) > 0.98
    id_token = any(token in n for token in ("uuid", "identifier", "identifiant", "code", "key", "numero", "number", "record_no", "record_number"))
    return named or (high_unique and id_token)


def _meaningful_columns(df: pd.DataFrame) -> tuple[list[str], list[str], list[str], list[str]]:
    rows = len(df)
    numeric: list[str] = []
    categorical: list[str] = []
    dates: list[str] = []
    identifiers: list[str] = []
    for col in df.columns:
        s = df[col]
        if _is_identifier_series(str(col), s, rows):
            identifiers.append(str(col)); continue
        if pd.api.types.is_numeric_dtype(s):
            numeric.append(str(col)); continue
        parsed = pd.to_datetime(s, errors="coerce", format="mixed")
        if len(s) and parsed.notna().mean() >= .8:
            dates.append(str(col)); continue
        unique = int(s.nunique(dropna=True))
        if 2 <= unique <= min(30, max(8, int(max(rows, 1) * .25))):
            categorical.append(str(col))
    return numeric, categorical, dates, identifiers


def _top_correlation(df: pd.DataFrame, numeric: list[str]) -> dict[str, Any] | None:
    if len(numeric) < 2:
        return None
    cols = numeric[:14]
    corr = df[cols].apply(pd.to_numeric, errors="coerce").corr(min_periods=3)
    best = None
    for i, a in enumerate(cols):
        for b in cols[i + 1:]:
            v = corr.loc[a, b]
            if pd.isna(v):
                continue
            row = {"x": a, "y": b, "coefficient": float(v), "abs": abs(float(v))}
            if best is None or row["abs"] > best["abs"]:
                best = row
    return best


def _figure_insight(viz: dict[str, Any], df: pd.DataFrame) -> str:
    typ = str(viz.get("type", "")).lower()
    data = viz.get("data") or []
    try:
        if typ == "bar" and data:
            top = max(data, key=lambda x: float(x.get("value", 0) or 0))
            return f"La valeur la plus elevee est observee pour {top.get('label')} ({_fmt(float(top.get('value', 0) or 0))})."
        if typ in {"line", "area"} and len(data) >= 2:
            first, last = float(data[0].get("value", 0) or 0), float(data[-1].get("value", 0) or 0)
            delta = last - first
            pct = (delta / abs(first) * 100) if first else None
            change = f"{pct:+.1f}%" if pct is not None and math.isfinite(pct) else _fmt(delta)
            return f"Entre le premier et le dernier point affiche, la mesure evolue de {change}."
        if typ == "histogram" and data:
            peak = max(data, key=lambda x: int(x.get("count", 0) or 0))
            return f"La classe la plus frequente se situe entre {_fmt(peak.get('from'))} et {_fmt(peak.get('to'))}."
        if typ == "scatter":
            x, y = viz.get("x"), viz.get("y")
            if x in df.columns and y in df.columns:
                pair = df[[x, y]].apply(pd.to_numeric, errors="coerce").dropna()
                if len(pair) >= 3:
                    r = float(pair.corr().iloc[0, 1])
                    return f"La correlation descriptive entre {x} et {y} est r = {r:.3f}; elle ne constitue pas une preuve de causalite."
        if typ == "heatmap":
            cols = viz.get("columns") or []
            matrix = {str(r.get("variable")): r for r in data if isinstance(r, dict)}
            best = None
            for i, a in enumerate(cols):
                for b in cols[i + 1:]:
                    v = matrix.get(str(a), {}).get(b)
                    if v is None: continue
                    row = (abs(float(v)), a, b, float(v))
                    if best is None or row[0] > best[0]: best = row
            if best:
                return f"La relation lineaire la plus forte de cette matrice concerne {best[1]} et {best[2]} (r = {best[3]:.3f})."
        if typ == "box" and data:
            meds = [(str(r.get("group")), float(r.get("median"))) for r in data if r.get("median") is not None]
            if len(meds) >= 2:
                lo, hi = min(meds, key=lambda x: x[1]), max(meds, key=lambda x: x[1])
                return f"Les medianes vont de {_fmt(lo[1])} ({lo[0]}) a {_fmt(hi[1])} ({hi[0]})."
    except Exception:
        pass
    return "Cette figure complete la lecture descriptive; les conclusions doivent etre confrontees aux tests et au contexte metier."


def _auto_visualizations(df: pd.DataFrame, meta: dict[str, Any], limit: int = 6) -> list[dict[str, Any]]:
    limit = max(1, min(int(limit), 10))
    numeric, categorical, dates, _ = _meaningful_columns(df)
    items: list[dict[str, Any]] = []

    def add(title: str, viz: dict[str, Any], reason: str):
        if len(items) >= limit:
            return
        signature = (viz.get("type"), viz.get("x"), viz.get("y"), viz.get("aggregation"))
        if any((x.get("visualization", {}).get("type"), x.get("visualization", {}).get("x"), x.get("visualization", {}).get("y"), x.get("visualization", {}).get("aggregation")) == signature for x in items):
            return
        items.append({
            "id": f"auto-{uuid4()}", "dataset_id": meta["id"], "dataset_version": int(meta.get("version", 1)),
            "title": title, "visualization": viz, "source": "auto_report", "reason": reason,
            "insight": _figure_insight(viz, df),
        })

    missing = df.isna().mean().sort_values(ascending=False)
    missing = missing[missing > 0].head(10)
    if len(missing):
        viz = {"type": "bar", "x": "variable", "y": "missing_pct", "aggregation": "none", "title": "Taux de valeurs manquantes", "data": [{"label": str(k), "value": float(v * 100)} for k, v in missing.items()]}
        add("Valeurs manquantes par variable", viz, "Prioriser les variables qui necessitent un traitement de qualite.")

    top_corr = _top_correlation(df, numeric)
    if len(numeric) >= 2:
        cols = numeric[:10]
        corr = df[cols].apply(pd.to_numeric, errors="coerce").corr(min_periods=3)
        matrix=[]
        for row in cols:
            item={"variable":row}
            for col in cols:
                val=corr.loc[row,col]; item[col]=None if pd.isna(val) else float(val)
            matrix.append(item)
        add("Carte des correlations", {"type":"heatmap","columns":cols,"data":matrix,"title":"Matrice de correlation — Pearson"}, "Identifier rapidement les associations lineaires entre mesures.")

    if top_corr and top_corr["abs"] >= .25:
        try:
            viz = build_visualization(df, chart_type="scatter", x=top_corr["x"], y=top_corr["y"])
            add(f"Relation entre {top_corr['x']} et {top_corr['y']}", viz, "Visualiser la relation lineaire la plus marquee parmi les variables numeriques.")
        except Exception:
            pass

    if numeric:
        candidates=[]
        for col in numeric[:12]:
            s=pd.to_numeric(df[col],errors="coerce").dropna()
            if len(s)>=5 and s.nunique()>2:
                cv=float(s.std(ddof=1)/(abs(s.mean())+1e-9)) if len(s)>1 else 0
                candidates.append((abs(cv),col))
        if candidates:
            col=max(candidates)[1]
            try:
                viz=build_visualization(df,chart_type="histogram",x=col,bins=20)
                add(f"Distribution de {col}",viz,"Examiner la forme, la dispersion et les zones de concentration de la variable.")
            except Exception:
                pass

    if categorical and numeric:
        cat = min(categorical[:8], key=lambda c: df[c].nunique(dropna=True))
        num = numeric[0]
        try:
            viz=build_visualization(df,chart_type="bar",x=cat,y=num,aggregation="mean")
            add(f"{num} moyen par {cat}",viz,"Comparer la mesure principale entre groupes sans confondre cette comparaison descriptive avec un test de significativite.")
        except Exception:
            pass
        if len(items)<limit:
            try:
                viz=build_visualization(df,chart_type="box",x=cat,y=num)
                add(f"Dispersion de {num} par {cat}",viz,"Comparer medianes, dispersion et valeurs extremes entre groupes.")
            except Exception:
                pass
    elif categorical:
        cat=categorical[0]
        try:
            viz=build_visualization(df,chart_type="bar",x=cat,aggregation="count")
            add(f"Repartition de {cat}",viz,"Visualiser le poids relatif des categories principales.")
        except Exception:
            pass

    if dates and numeric and len(items)<limit:
        try:
            viz=build_visualization(df,chart_type="line",x=dates[0],y=numeric[0],aggregation="mean")
            add(f"Evolution de {numeric[0]} dans le temps",viz,"Reperer les tendances ou changements temporels visibles dans les donnees.")
        except Exception:
            pass
    return items[:limit]


def _analytical_story(df: pd.DataFrame, profile: dict[str, Any], quality: dict[str, Any], decision: dict[str, Any], analysis: dict[str, Any] | None) -> dict[str, Any]:
    numeric, categorical, dates, identifiers = _meaningful_columns(df)
    rows = int(profile.get("rows", len(df)))
    findings: list[dict[str, Any]] = []
    qscore = int(quality.get("score", 0))
    findings.append({
        "rank": 1, "severity": "critical" if qscore < 70 else "attention" if qscore < 90 else "positive",
        "title": "Fiabilite des donnees",
        "statement": f"Le score de qualite est de {qscore}/100 avec {quality.get('issues_count', 0)} alerte(s) detectee(s).",
        "evidence": "Moteur Data Quality",
        "interpretation": "Les analyses subsequentes doivent etre lues a la lumiere de ces controles de qualite." if quality.get("issues_count") else "Aucune alerte majeure n'a ete detectee par les regles actuellement actives.",
    })
    miss=df.isna().mean().sort_values(ascending=False)
    if len(miss) and float(miss.iloc[0])>0:
        col=str(miss.index[0]); pct=float(miss.iloc[0]*100)
        findings.append({"rank":2,"severity":"critical" if pct>=30 else "attention","title":"Completude","statement":f"{col} est la variable la plus incomplete avec {pct:.1f}% de valeurs manquantes.","evidence":"Calcul de completude par colonne","interpretation":"Une strategie d'imputation ou une analyse de sensibilite peut etre necessaire avant modelisation."})
    corr=_top_correlation(df,numeric)
    if corr and corr["abs"]>=.3:
        findings.append({"rank":3,"severity":"info","title":"Association lineaire principale","statement":f"{corr['x']} et {corr['y']} presentent la correlation lineaire la plus forte (r = {corr['coefficient']:.3f}).","evidence":"Correlation de Pearson sur observations completes","interpretation":"Cette association merite une analyse bivariée ou multivariee; elle ne demontre pas une causalite."})
    skewed=[]
    for col in numeric[:16]:
        ser=pd.to_numeric(df[col],errors="coerce").dropna()
        if len(ser)>=8 and ser.nunique()>2:
            sk=float(ser.skew())
            if math.isfinite(sk): skewed.append((abs(sk),col,sk))
    if skewed:
        _,col,sk=max(skewed)
        if abs(sk)>=1:
            findings.append({"rank":4,"severity":"attention","title":"Distribution asymetrique","statement":f"{col} presente une asymetrie marquee (skewness = {sk:.2f}).","evidence":"Coefficient d'asymetrie","interpretation":"La mediane, les quantiles ou une transformation peuvent etre plus informatifs qu'une moyenne seule."})
    if categorical:
        cat=categorical[0]; vc=df[cat].astype("string").value_counts(dropna=True)
        if len(vc):
            share=float(vc.iloc[0]/max(vc.sum(),1)*100)
            if share>=50:
                findings.append({"rank":5,"severity":"info","title":"Concentration categorielle","statement":f"La modalite {vc.index[0]} represente {share:.1f}% des observations renseignees de {cat}.","evidence":"Frequences par categorie","interpretation":"Cette concentration peut influencer les comparaisons entre groupes et la validation des modeles."})
    if dates and numeric:
        date_col, num_col=dates[0],numeric[0]
        work=pd.DataFrame({"d":pd.to_datetime(df[date_col],errors="coerce",format="mixed"),"v":pd.to_numeric(df[num_col],errors="coerce")}).dropna().sort_values("d")
        if len(work)>=6:
            n=max(2,len(work)//3); first=float(work.head(n)["v"].mean()); last=float(work.tail(n)["v"].mean()); delta=last-first
            if abs(delta)>1e-9:
                findings.append({"rank":6,"severity":"info","title":"Evolution temporelle descriptive","statement":f"La moyenne de {num_col} passe de {_fmt(first)} au debut de la periode a {_fmt(last)} en fin de periode.","evidence":f"Comparaison du premier et du dernier tiers selon {date_col}","interpretation":"Ce constat decrit une evolution observee; un modele de serie temporelle est necessaire pour quantifier tendance et incertitude."})
    if analysis:
        for ix,f in enumerate(analysis.get("findings",[])[:2],start=7):
            findings.append({"rank":ix,"severity":str(f.get("level","info")),"title":str(f.get("title","Constat AI Analyst")),"statement":str(f.get("statement","")),"evidence":f"AI Analyst / {f.get('evidence',{}).get('tool','outil execute')}","interpretation":"Constat issu d'un moteur analytique execute et conserve avec sa provenance."})
    opening=f"L'analyse porte sur {rows} observations et {int(profile.get('columns_count', len(df.columns)))} variables. Les resultats ci-dessous sont classes par importance analytique et relies a des calculs deterministes."
    conclusion="Les principaux signaux doivent etre combines avec le contexte metier, les hypotheses statistiques et les limites du dataset avant toute decision."
    return {"opening":opening,"findings":findings[:8],"conclusion":conclusion,"identifiers_excluded":identifiers,"calculation_policy":"deterministic_engines_only"}


def _report_limitations(df: pd.DataFrame, profile: dict[str, Any], quality: dict[str, Any], analysis: dict[str, Any] | None) -> list[dict[str, str]]:
    numeric, categorical, dates, identifiers = _meaningful_columns(df)
    rows=len(df); total=max(rows*max(len(df.columns),1),1); missing_pct=float(df.isna().sum().sum()/total*100)
    items=[]
    if missing_pct>0:
        items.append({"title":"Valeurs manquantes","text":f"{missing_pct:.1f}% des cellules sont manquantes.","mitigation":"Documenter l'imputation, comparer plusieurs strategies et verifier la robustesse des conclusions."})
    if quality.get("issues_count",0):
        items.append({"title":"Qualite des donnees","text":f"{quality.get('issues_count')} alerte(s) de qualite restent signalee(s).","mitigation":"Traiter ou justifier ces alertes avant de diffuser des conclusions operationnelles."})
    if rows<100:
        items.append({"title":"Taille d'echantillon","text":f"Le dataset ne contient que {rows} observations.","mitigation":"Limiter la complexite des modeles et utiliser des intervalles d'incertitude / validation adaptee."})
    if identifiers:
        items.append({"title":"Identifiants exclus des analyses automatiques","text":", ".join(identifiers[:8]),"mitigation":"Ces colonnes restent disponibles mais ne sont pas interpretees comme variables explicatives par defaut."})
    if not dates:
        items.append({"title":"Temporalite non etablie","text":"Aucune variable temporelle fiable n'a ete detectee automatiquement.","mitigation":"Ne pas interpreter les associations comme une evolution dans le temps sans variable date valide."})
    items.append({"title":"Causalite","text":"Les associations, correlations et differences descriptives ne prouvent pas un lien causal.","mitigation":"Utiliser un protocole causal ou experimental lorsque la question porte sur un effet causal."})
    if not analysis:
        items.append({"title":"Analyse contextuelle","text":"Aucune session AI Analyst n'est liee a ce rapport.","mitigation":"Associer une session analytique lorsque le rapport doit documenter une question metier precise."})
    return items[:7]


def build_report(
    dataset_id: str,
    title: str,
    sections: list[str] | None = None,
    analysis_session_id: str | None = None,
    *,
    template: str = "analytical",
    subtitle: str | None = None,
    author: str | None = None,
    organization: str | None = None,
    visualization_ids: list[str] | None = None,
    auto_story: bool = False,
    auto_visualizations: bool = False,
    max_visualizations: int = 6,
) -> dict[str, Any]:
    meta = get_meta(dataset_id)
    df = load_dataframe(dataset_id)
    profile = profile_dataframe(df)
    quality = quality_report(df)
    decision = decision_support(profile, quality)
    template = template if template in TEMPLATES else "analytical"
    allowed = set(SECTION_LABELS)
    defaults = TEMPLATES[template]["default_sections"]
    chosen = [s for s in (sections or defaults) if s in allowed]
    analysis = None
    if analysis_session_id:
        analysis = get_analysis(analysis_session_id)
        if analysis.get("provenance", {}).get("dataset_id") != dataset_id:
            raise ValueError("L'analyse selectionnee n'appartient pas a ce dataset")

    max_visualizations = max(1, min(int(max_visualizations), 10))
    saved_all = [v for v in list_visualizations(dataset_id) if int(v.get("dataset_version", 0)) == int(meta.get("version", 1))]
    if visualization_ids:
        wanted = set(visualization_ids)
        saved = [v for v in saved_all if v.get("id") in wanted][:max_visualizations]
    else:
        saved = saved_all[:max_visualizations]
    auto_items: list[dict[str, Any]] = []
    if auto_visualizations and len(saved) < max_visualizations:
        auto_items = _auto_visualizations(df, meta, max_visualizations - len(saved))
        existing = {(v.get("visualization", {}).get("type"), v.get("visualization", {}).get("x"), v.get("visualization", {}).get("y")) for v in saved}
        auto_items = [v for v in auto_items if (v.get("visualization", {}).get("type"), v.get("visualization", {}).get("x"), v.get("visualization", {}).get("y")) not in existing]
        saved = (saved + auto_items)[:max_visualizations]

    blocks: list[dict[str, Any]] = []
    if "executive_summary" in chosen:
        blocks.append({"type": "executive_summary", "title": SECTION_LABELS["executive_summary"], "data": _executive_summary(profile, quality, decision, analysis)})
    if "analytical_story" in chosen:
        blocks.append({"type": "analytical_story", "title": SECTION_LABELS["analytical_story"], "data": _analytical_story(df, profile, quality, decision, analysis)})
    if "overview" in chosen:
        blocks.append({"type": "overview", "title": SECTION_LABELS["overview"], "data": {
            "rows": profile.get("rows"), "columns": profile.get("columns_count"), "duplicates": profile.get("duplicates"),
            "memory_bytes": profile.get("memory_bytes"), "quality_score": quality.get("score"),
            "type_summary": _type_summary(profile), "missing_summary": _missing_summary(profile),
            "recommended_actions": decision.get("actions", [])[:5],
        }})
    if "quality" in chosen:
        blocks.append({"type": "quality", "title": SECTION_LABELS["quality"], "data": {
            "score": quality.get("score"), "issues_count": quality.get("issues_count"), "severity": _severity_counts(quality),
            "issues": quality.get("issues", [])[:30]
        }})
    if "descriptive" in chosen:
        blocks.append({"type": "table", "title": SECTION_LABELS["descriptive"], "columns": ["variable", "n", "mean", "median", "std", "q1", "q3", "min", "max"], "rows": _numeric_summary(df)})
    if "visualizations" in chosen:
        blocks.append({"type": "visualizations", "title": SECTION_LABELS["visualizations"], "items": saved})
    if "ai_analysis" in chosen:
        if analysis:
            blocks.append({"type": "ai_analysis", "title": SECTION_LABELS["ai_analysis"], "data": {
                "session_id": analysis.get("session_id"), "question": analysis.get("question"), "answer": analysis.get("answer"),
                "intent": analysis.get("intent"), "critic": analysis.get("critic"), "findings": analysis.get("findings", []),
                "provenance": analysis.get("provenance", {}),
            }})
        else:
            blocks.append({"type": "note", "title": SECTION_LABELS["ai_analysis"], "text": "Aucune session AI Analyst n'a ete selectionnee pour ce rapport."})
    if "limitations" in chosen:
        blocks.append({"type": "limitations", "title": SECTION_LABELS["limitations"], "items": _report_limitations(df, profile, quality, analysis)})
    if "methodology" in chosen:
        blocks.append({"type": "methodology", "title": SECTION_LABELS["methodology"], "items": [
            "Les statistiques et metriques sont calculees par les moteurs executables de DataVision, et non generees par un LLM.",
            "Le rapport reference la version exacte du dataset utilisee au moment de sa creation.",
            "Les donnees originales restent immuables; les transformations sont versionnees et tracables.",
            "Les tests statistiques doivent etre interpretes avec leurs hypotheses, tailles d'effet et limites.",
            "Les resultats predicitifs doivent etre lus avec leurs metriques de validation et le contexte metier.",
        ]})
    if "provenance" in chosen:
        blocks.append({"type": "provenance", "title": SECTION_LABELS["provenance"], "data": {
            "dataset_id": meta["id"], "dataset_name": meta.get("original_name"), "dataset_version": meta.get("version", 1),
            "root_id": meta.get("root_id", meta["id"]), "parent_id": meta.get("parent_id"), "operation": meta.get("operation"),
            "generated_at": _now(), "analysis_session_id": analysis_session_id,
        }})

    report = {
        "id": str(uuid4()), "dataset_id": dataset_id,
        "title": title.strip() or f"Rapport - {meta.get('original_name', 'Dataset')}",
        "subtitle": (subtitle or "Analyse de donnees et aide a la decision").strip(),
        "author": (author or "DataVision AI").strip(), "organization": (organization or "").strip(),
        "template": template, "template_label": TEMPLATES[template]["label"],
        "created_at": _now(), "dataset": {"id": meta["id"], "name": meta.get("original_name"), "version": meta.get("version", 1)},
        "sections": chosen, "section_outline": [{"number": i + 1, "key": key, "title": SECTION_LABELS[key]} for i, key in enumerate(chosen)],
        "analysis_session_id": analysis_session_id,
        "visualization_ids": [v.get("id") for v in saved if v.get("source") != "auto_report"],
        "auto_visualization_count": len([v for v in saved if v.get("source") == "auto_report"]),
        "generation_mode": "intelligent" if (auto_story or auto_visualizations) else "manual",
        "intelligence": {"auto_story": bool(auto_story), "auto_visualizations": bool(auto_visualizations), "max_visualizations": max_visualizations},
        "blocks": blocks,
        "reproducibility": {"dataset_version_locked": True, "analysis_session_locked": bool(analysis_session_id), "visualizations_locked": True, "auto_generated_visualizations_embedded": bool(auto_items)},
    }
    (_dir() / f"{report['id']}.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return report


def get_report(report_id: str) -> dict[str, Any]:
    path = _dir() / f"{report_id}.json"
    if not path.exists():
        raise FileNotFoundError(report_id)
    return json.loads(path.read_text(encoding="utf-8"))


def list_reports(dataset_id: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for path in _dir().glob("*.json"):
        try:
            report = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if report.get("dataset_id") != dataset_id:
            continue
        rows.append({
            "id": report.get("id"), "title": report.get("title"), "created_at": report.get("created_at"),
            "dataset_version": report.get("dataset", {}).get("version"), "sections": report.get("sections", []),
            "analysis_session_id": report.get("analysis_session_id"), "template": report.get("template", "analytical"),
        })
    rows.sort(key=lambda x: x.get("created_at") or "", reverse=True)
    return rows


# ----------------------------- visual renderer -----------------------------

def _color(value: str):
    return colors.HexColor(value)


def _chart_drawing(item: dict[str, Any], width: float = 480, height: float = 245) -> Drawing | None:
    viz = item.get("visualization", {}) if "visualization" in item else item
    typ = str(viz.get("type", "")).lower()
    data = viz.get("data") or []
    if not isinstance(data, list) or not data:
        return None
    d = Drawing(width, height)
    left, right, top, bottom = 52, 18, 24, 42
    pw, ph = width - left - right, height - top - bottom
    d.add(Rect(0, 0, width, height, fillColor=_color("#FFFFFF"), strokeColor=_color(BRAND["line"]), strokeWidth=.6, rx=6, ry=6))
    d.add(Line(left, bottom, left, bottom + ph, strokeColor=_color("#B9C9D1"), strokeWidth=.7))
    d.add(Line(left, bottom, left + pw, bottom, strokeColor=_color("#B9C9D1"), strokeWidth=.7))

    def label(text: Any, x: float, y: float, size: int = 7, anchor: str = "middle", color: str = BRAND["muted"]):
        s = str(text)
        if len(s) > 16:
            s = s[:14] + ".."
        d.add(String(x, y, s, fontName="Helvetica", fontSize=size, textAnchor=anchor, fillColor=_color(color)))

    def y_grid(maxv: float, minv: float = 0.0):
        if math.isclose(maxv, minv):
            maxv = minv + 1
        for i in range(5):
            frac = i / 4
            y = bottom + frac * ph
            val = minv + frac * (maxv - minv)
            d.add(Line(left, y, left + pw, y, strokeColor=_color("#EDF2F4"), strokeWidth=.4))
            label(_fmt(val), left - 7, y - 2, 6, "end")

    palette = ["#2798B7", "#5CB8CF", "#1D708A", "#78C8A8", "#E0AE57", "#B86D79"]

    if typ in {"bar", "line", "area"}:
        vals = [float(x.get("value", 0) or 0) for x in data[:80]]
        if not vals:
            return None
        minv = min(0.0, min(vals)); maxv = max(0.0, max(vals)); span = max(maxv - minv, 1e-9)
        y_grid(maxv, minv)
        n = len(vals)
        if typ == "bar":
            gap = max(2, pw / max(n, 1) * .18); bw = max(2, (pw - gap * (n + 1)) / max(n, 1))
            for i, (row, val) in enumerate(zip(data[:80], vals)):
                x0 = left + gap + i * (bw + gap); y0 = bottom + (0 - minv) / span * ph; yv = bottom + (val - minv) / span * ph
                d.add(Rect(x0, min(y0, yv), bw, max(abs(yv - y0), 1), fillColor=_color(palette[i % len(palette)]), strokeColor=None, rx=1, ry=1))
                if n <= 14 or i % max(1, n // 10) == 0:
                    label(row.get("label", i + 1), x0 + bw / 2, bottom - 12, 6)
        else:
            pts = []
            for i, val in enumerate(vals):
                x = left + (i / max(n - 1, 1)) * pw; y = bottom + (val - minv) / span * ph; pts.append((x, y))
            if typ == "area" and len(pts) >= 2:
                base = bottom + (0 - minv) / span * ph
                poly = [pts[0][0], base]
                for x, y in pts: poly.extend([x, y])
                poly.extend([pts[-1][0], base])
                d.add(Polygon(poly, fillColor=_color("#D9F0F6"), strokeColor=None))
            for i in range(1, len(pts)):
                d.add(Line(pts[i-1][0], pts[i-1][1], pts[i][0], pts[i][1], strokeColor=_color(BRAND["teal"]), strokeWidth=1.6))
            for i, (x, y) in enumerate(pts):
                if n <= 60: d.add(Circle(x, y, 1.8, fillColor=_color(BRAND["teal_dark"]), strokeColor=None))
                if n <= 12 or i % max(1, n // 8) == 0: label(data[i].get("label", i + 1), x, bottom - 12, 6)
        return d

    if typ == "histogram":
        vals = [float(x.get("count", 0) or 0) for x in data]
        maxv = max(vals) if vals else 1; y_grid(maxv)
        n = len(vals); bw = pw / max(n, 1)
        for i, (row, val) in enumerate(zip(data, vals)):
            h = ph * val / max(maxv, 1e-9)
            d.add(Rect(left + i * bw + .5, bottom, max(bw - 1, 1), h, fillColor=_color("#75C1D4"), strokeColor=_color("#3B91AA"), strokeWidth=.2))
        label(viz.get("x", "Valeur"), left + pw / 2, 9, 7)
        return d

    if typ == "density":
        xs = [float(x.get("x", 0) or 0) for x in data]; ys = [float(x.get("density", 0) or 0) for x in data]
        if not xs or not ys: return None
        xmin, xmax = min(xs), max(xs); ymin, ymax = 0.0, max(ys); y_grid(ymax)
        pts = [(left + (x - xmin) / max(xmax - xmin, 1e-9) * pw, bottom + (y - ymin) / max(ymax - ymin, 1e-9) * ph) for x, y in zip(xs, ys)]
        for i in range(1, len(pts)):
            d.add(Line(*pts[i-1], *pts[i], strokeColor=_color(BRAND["teal"]), strokeWidth=1.5))
        label(_fmt(xmin), left, bottom - 12, 6); label(_fmt(xmax), left + pw, bottom - 12, 6)
        return d

    if typ == "scatter":
        pts_raw = data[:800]
        xs = [float(x.get("x", 0) or 0) for x in pts_raw]; ys = [float(x.get("y", 0) or 0) for x in pts_raw]
        if not xs or not ys: return None
        xmin, xmax = min(xs), max(xs); ymin, ymax = min(ys), max(ys); y_grid(ymax, ymin)
        for row, x, y in zip(pts_raw, xs, ys):
            px = left + (x - xmin) / max(xmax - xmin, 1e-9) * pw; py = bottom + (y - ymin) / max(ymax - ymin, 1e-9) * ph
            d.add(Circle(px, py, 1.6, fillColor=_color("#2798B7"), strokeColor=None, fillOpacity=.55))
        label(viz.get("x", "X"), left + pw / 2, 9, 7); label(viz.get("y", "Y"), 12, bottom + ph / 2, 7)
        return d

    if typ == "heatmap":
        columns = viz.get("columns") or []
        columns = columns[:12]
        if not columns: return None
        n = len(columns); cell = min(pw / n, ph / n); startx = left + max((pw - cell * n) / 2, 0); starty = bottom + max((ph - cell * n) / 2, 0)
        matrix = {str(r.get("variable")): r for r in data if isinstance(r, dict)}
        for i, row_name in enumerate(columns):
            row = matrix.get(str(row_name), {})
            for j, col_name in enumerate(columns):
                val = row.get(col_name)
                v = 0.0 if val is None else max(-1.0, min(1.0, float(val)))
                if v >= 0:
                    base = colors.Color(1 - .65*v, 1 - .38*v, 1 - .28*v)
                else:
                    a = abs(v); base = colors.Color(1 - .28*a, 1 - .50*a, 1 - .35*a)
                x = startx + j * cell; y = starty + (n - i - 1) * cell
                d.add(Rect(x, y, cell, cell, fillColor=base, strokeColor=colors.white, strokeWidth=.5))
                if n <= 8: label(_fmt(v), x + cell/2, y + cell/2 - 2, 5, "middle", BRAND["ink"])
            label(row_name, startx - 5, starty + (n - i - .5) * cell - 2, 5, "end")
        for j, col_name in enumerate(columns): label(col_name, startx + (j + .5) * cell, starty - 11, 5)
        return d

    if typ == "box":
        groups = data[:16]
        all_vals = []
        for r in groups:
            for key in ["whisker_low", "q1", "median", "q3", "whisker_high"]:
                if r.get(key) is not None: all_vals.append(float(r[key]))
        if not all_vals: return None
        ymin, ymax = min(all_vals), max(all_vals); y_grid(ymax, ymin)
        n = len(groups); slot = pw / max(n, 1)
        for i, r in enumerate(groups):
            cx = left + (i + .5) * slot; bw = min(28, slot * .55)
            def py(v): return bottom + (float(v) - ymin) / max(ymax-ymin,1e-9) * ph
            q1, med, q3 = py(r["q1"]), py(r["median"]), py(r["q3"]); lo, hi = py(r["whisker_low"]), py(r["whisker_high"])
            d.add(Line(cx, lo, cx, hi, strokeColor=_color("#4F6774"), strokeWidth=.7)); d.add(Rect(cx-bw/2, q1, bw, max(q3-q1,1), fillColor=_color("#D9F0F6"), strokeColor=_color(BRAND["teal"]), strokeWidth=1)); d.add(Line(cx-bw/2, med, cx+bw/2, med, strokeColor=_color(BRAND["navy"]), strokeWidth=1.2)); d.add(Line(cx-bw*.28, lo, cx+bw*.28, lo, strokeColor=_color("#4F6774"), strokeWidth=.7)); d.add(Line(cx-bw*.28, hi, cx+bw*.28, hi, strokeColor=_color("#4F6774"), strokeWidth=.7))
            label(r.get("group", i + 1), cx, bottom - 12, 6)
        return d
    return None


def _drawing_svg(item: dict[str, Any]) -> str:
    drawing = _chart_drawing(item, 640, 320)
    if drawing is None:
        return ""
    raw = renderSVG.drawToString(drawing)
    return raw.decode("utf-8") if isinstance(raw, bytes) else str(raw)


def _drawing_png(item: dict[str, Any], path: Path) -> bool:
    drawing = _chart_drawing(item, 700, 350)
    if drawing is None:
        return False
    renderPM.drawToFile(drawing, str(path), fmt="PNG", dpi=120)
    return path.exists() and path.stat().st_size > 0


# ----------------------------- markdown / html -----------------------------

def _markdown(report: dict[str, Any]) -> str:
    out = [
        f"# {report['title']}", "", report.get("subtitle", ""), "",
        f"**Dataset :** {report['dataset']['name']} - version {report['dataset']['version']}",
        f"**Auteur :** {report.get('author') or 'DataVision AI'}",
        f"**Genere :** {_format_date(report.get('created_at'))}", "",
        "## Sommaire", "",
    ]
    for entry in report.get("section_outline", []):
        out.append(f"{entry['number']}. {entry['title']}")
    out.append("")
    section_number = 0
    for block in report.get("blocks", []):
        section_number += 1
        out.extend([f"## {section_number}. {block.get('title','Section')}", ""])
        kind = block.get("type")
        if kind == "executive_summary":
            d = block["data"]
            out.append("### Indicateurs cles")
            for kpi in d.get("kpis", []): out.append(f"- **{kpi['label']} :** {_fmt(kpi['value'])} {kpi.get('hint','')}")
            out.append("\n### Points cles")
            for h in d.get("highlights", []): out.append(f"- **{h.get('title')}** - {h.get('text')}")
            out.append("\n### Recommandations")
            for r in d.get("recommendations", []): out.append(f"- **{r.get('priority','').upper()}** - {r.get('action','')} ({r.get('evidence','')})")
            out.append("")
        elif kind == "analytical_story":
            d = block["data"]
            out.extend([d.get("opening", ""), "", "### Constats prioritaires"])
            for f in d.get("findings", []):
                out.append(f"- **{f.get('title','')}** — {f.get('statement','')}  \n  *Preuve :* {f.get('evidence','')}  \n  *Lecture :* {f.get('interpretation','')}")
            out.extend(["", "### Conclusion de lecture", d.get("conclusion", ""), ""])
        elif kind == "overview":
            d = block["data"]
            out.extend([f"- Lignes : {_fmt(d.get('rows'))}", f"- Variables : {_fmt(d.get('columns'))}", f"- Doublons : {_fmt(d.get('duplicates'))}", f"- Score qualite : {_fmt(d.get('quality_score'))}/100", ""])
        elif kind == "quality":
            d = block["data"]
            out.append(f"Score qualite : **{_fmt(d.get('score'))}/100** - {d.get('issues_count',0)} probleme(s).\n")
            for issue in d.get("issues", []): out.append(f"- **{str(issue.get('severity','info')).upper()}** - {issue.get('column','General')} - {issue.get('description','')} Recommandation : {issue.get('recommendation','')}")
            out.append("")
        elif kind == "table":
            cols = block.get("columns", [])
            out.append("| " + " | ".join(cols) + " |"); out.append("| " + " | ".join(["---"] * len(cols)) + " |")
            for row in block.get("rows", []): out.append("| " + " | ".join(_fmt(row.get(c)) for c in cols) + " |")
            out.append("")
        elif kind == "visualizations":
            items = block.get("items", [])
            if not items: out.append("Aucune visualisation epinglee pour cette version du dataset.\n")
            for idx, item in enumerate(items, 1):
                viz = item.get("visualization", {})
                out.append(f"### Figure {idx}. {item.get('title','Visualisation')}")
                out.append(f"Type : `{viz.get('type','-')}` - dataset version {item.get('dataset_version','-')}")
                if item.get("reason"): out.append(f"**Pourquoi cette figure :** {item.get('reason')}")
                if item.get("insight"): out.append(f"**Lecture :** {item.get('insight')}")
                out.append("")
        elif kind == "ai_analysis":
            d = block["data"]
            out.extend([f"**Question :** {d.get('question','')}", "", d.get("answer", ""), "", "### Constats"])
            for f in d.get("findings", []): out.append(f"- **{f.get('title','')}** - {f.get('statement','')}")
            out.append("")
        elif kind == "limitations":
            for item in block.get("items", []):
                out.append(f"- **{item.get('title','')}** — {item.get('text','')}  \n  *Mesure de prudence :* {item.get('mitigation','')}")
            out.append("")
        elif kind == "methodology":
            for item in block.get("items", []): out.append(f"- {item}")
            out.append("")
        elif kind == "provenance":
            for k, v in block.get("data", {}).items(): out.append(f"- {k}: `{_fmt(v)}`")
            out.append("")
        else:
            out.extend([block.get("text", ""), ""])
    return "\n".join(out)


def _html_table(columns: list[str], rows: list[dict[str, Any]], limit: int | None = None) -> str:
    if limit is not None: rows = rows[:limit]
    head = "".join(f"<th>{html.escape(str(c))}</th>" for c in columns)
    body = "".join("<tr>" + "".join(f"<td>{html.escape(_fmt(r.get(c)))}</td>" for c in columns) + "</tr>" for r in rows)
    return f"<div class='table-wrap'><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>"


def _html(report: dict[str, Any]) -> str:
    toc = "".join(f"<a href='#sec-{e['number']}'><span>{e['number']:02d}</span>{html.escape(e['title'])}</a>" for e in report.get("section_outline", []))
    sections: list[str] = []
    sec = 0
    for block in report.get("blocks", []):
        sec += 1
        kind = block.get("type")
        content = ""
        if kind == "executive_summary":
            d = block["data"]
            kpis = "".join(f"<div class='kpi'><span>{html.escape(k['label'])}</span><b>{html.escape(_fmt(k['value']))}</b><small>{html.escape(k.get('hint',''))}</small></div>" for k in d.get("kpis", []))
            highlights = "".join(f"<article class='insight {html.escape(str(x.get('level','info')))}'><b>{html.escape(x.get('title',''))}</b><p>{html.escape(x.get('text',''))}</p></article>" for x in d.get("highlights", []))
            recs = "".join(f"<li><span class='priority'>{html.escape(str(x.get('priority','medium')).upper())}</span><div><b>{html.escape(x.get('action',''))}</b><small>{html.escape(x.get('evidence',''))}</small></div></li>" for x in d.get("recommendations", []))
            content = f"<div class='kpi-grid'>{kpis}</div><h3>Points cles</h3><div class='insight-grid'>{highlights}</div><h3>Recommandations</h3><ol class='recs'>{recs}</ol>"
        elif kind == "analytical_story":
            d = block["data"]
            findings = "".join(f"<article class='story-card {html.escape(str(x.get('severity','info')))}'><div class='story-rank'>{int(x.get('rank',0)):02d}</div><div><b>{html.escape(x.get('title',''))}</b><p>{html.escape(x.get('statement',''))}</p><small><strong>Preuve :</strong> {html.escape(x.get('evidence',''))}</small><small><strong>Lecture :</strong> {html.escape(x.get('interpretation',''))}</small></div></article>" for x in d.get("findings", []))
            content = f"<p class='lead'>{html.escape(d.get('opening',''))}</p><div class='story-grid'>{findings}</div><div class='story-conclusion'><span>CONCLUSION DE LECTURE</span><p>{html.escape(d.get('conclusion',''))}</p></div>"
        elif kind == "overview":
            d = block["data"]
            cards = [
                ("Observations", d.get("rows"), "lignes"), ("Variables", d.get("columns"), "colonnes"),
                ("Doublons", d.get("duplicates"), "lignes"), ("Qualite", d.get("quality_score"), "/100")
            ]
            content = "<div class='kpi-grid'>" + "".join(f"<div class='kpi'><span>{a}</span><b>{html.escape(_fmt(b))}</b><small>{c}</small></div>" for a,b,c in cards) + "</div>"
            if d.get("type_summary"): content += "<h3>Structure des variables</h3>" + _html_table(["type","count"], d["type_summary"])
            if d.get("missing_summary"): content += "<h3>Variables les plus incompletes</h3>" + _html_table(["variable","missing","missing_pct"], d["missing_summary"])
        elif kind == "quality":
            d = block["data"]
            content = f"<div class='score-ring'><b>{html.escape(_fmt(d.get('score')))}</b><span>/100</span><small>Score qualite</small></div>"
            rows = [{"severity": i.get("severity"), "column": i.get("column", "General"), "description": i.get("description"), "impact": i.get("impact"), "recommendation": i.get("recommendation")} for i in d.get("issues", [])]
            content += _html_table(["severity","column","description","impact","recommendation"], rows)
        elif kind == "table":
            content = _html_table(block.get("columns", []), block.get("rows", []))
        elif kind == "visualizations":
            cards=[]
            for idx,item in enumerate(block.get("items",[]),1):
                svg=_drawing_svg(item)
                viz=item.get("visualization",{})
                narrative = ""
                if item.get("reason"): narrative += f"<p class='figure-reason'><strong>Objectif :</strong> {html.escape(str(item.get('reason')))}</p>"
                if item.get("insight"): narrative += f"<p class='figure-insight'><strong>Lecture :</strong> {html.escape(str(item.get('insight')))}</p>"
                cards.append(f"<figure><div class='figure-head'><span>Figure {idx}</span><b>{html.escape(item.get('title','Visualisation'))}</b></div>{svg if svg else '<div class=\"empty-chart\">Apercu graphique indisponible</div>'}{narrative}<figcaption>{html.escape(str(viz.get('type','-')))} - dataset v{html.escape(str(item.get('dataset_version','-')))}</figcaption></figure>")
            content="<div class='figure-grid'>"+"".join(cards)+"</div>" if cards else "<p class='empty'>Aucune visualisation epinglee pour cette version.</p>"
        elif kind == "ai_analysis":
            d=block["data"]
            findings="".join(f"<article class='finding'><span>{html.escape(str(f.get('level','info')).upper())}</span><div><b>{html.escape(f.get('title',''))}</b><p>{html.escape(f.get('statement',''))}</p></div></article>" for f in d.get("findings",[]))
            content=f"<div class='question'><span>Question analytique</span><b>{html.escape(d.get('question',''))}</b></div><div class='answer'>{html.escape(d.get('answer',''))}</div><h3>Constats verifies</h3><div class='findings'>{findings}</div>"
        elif kind == "limitations":
            content = "<div class='limit-list'>" + "".join(f"<article><b>{html.escape(x.get('title',''))}</b><p>{html.escape(x.get('text',''))}</p><small><strong>Prudence :</strong> {html.escape(x.get('mitigation',''))}</small></article>" for x in block.get("items", [])) + "</div>"
        elif kind == "methodology":
            content="<div class='method-list'>"+"".join(f"<div><span>{i:02d}</span><p>{html.escape(x)}</p></div>" for i,x in enumerate(block.get("items",[]),1))+"</div>"
        elif kind == "provenance":
            content="<div class='provenance'>"+"".join(f"<div><span>{html.escape(str(k))}</span><b>{html.escape(_fmt(v))}</b></div>" for k,v in block.get("data",{}).items())+"</div>"
        else:
            content=f"<p>{html.escape(block.get('text',''))}</p>"
        sections.append(f"<section id='sec-{sec}'><header class='section-title'><span>{sec:02d}</span><div><small>SECTION</small><h2>{html.escape(block.get('title','Section'))}</h2></div></header>{content}</section>")

    css = f"""
    :root{{--navy:{BRAND['navy']};--teal:{BRAND['teal']};--ink:{BRAND['ink']};--muted:{BRAND['muted']};--line:{BRAND['line']};--surface:{BRAND['surface']};}}
    *{{box-sizing:border-box}} body{{margin:0;font-family:Inter,Segoe UI,Arial,sans-serif;color:var(--ink);background:#eef3f5;line-height:1.45}}
    .document{{max-width:1120px;margin:32px auto;background:white;box-shadow:0 18px 50px rgba(25,55,68,.12)}}
    .cover{{min-height:760px;padding:74px 72px 56px;position:relative;background:linear-gradient(155deg,#fff 0%,#fff 65%,#edf8fb 100%);display:flex;flex-direction:column}}
    .brand{{display:flex;align-items:center;gap:12px;color:var(--navy);font-weight:800}} .brand-mark{{width:42px;height:42px;border-radius:10px;background:var(--teal);color:white;display:grid;place-items:center}}
    .cover-main{{margin:auto 0;max-width:780px}} .cover-kicker{{text-transform:uppercase;letter-spacing:.18em;color:var(--teal);font-size:12px;font-weight:800}}
    h1{{font-size:48px;line-height:1.05;margin:16px 0 14px;color:var(--navy);letter-spacing:-.03em}} .subtitle{{font-size:21px;color:#5e7480;max-width:690px}}
    .cover-meta{{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:1px;background:var(--line);border:1px solid var(--line);margin-top:42px}} .cover-meta div{{background:#fff;padding:18px}} .cover-meta span{{display:block;color:var(--muted);font-size:11px;text-transform:uppercase;letter-spacing:.08em}} .cover-meta b{{display:block;margin-top:5px;font-size:14px}}
    .cover-foot{{display:flex;justify-content:space-between;border-top:1px solid var(--line);padding-top:18px;color:var(--muted);font-size:11px}}
    .toc{{padding:54px 72px;border-top:8px solid var(--teal)}} .toc h2{{font-size:28px;color:var(--navy)}} .toc-grid{{display:grid;grid-template-columns:1fr 1fr;gap:8px 28px}} .toc a{{text-decoration:none;color:var(--ink);display:flex;gap:12px;padding:11px 0;border-bottom:1px solid #edf2f4}} .toc a span{{color:var(--teal);font-weight:800}}
    section{{padding:52px 72px;border-top:1px solid #e8eef1;break-before:page}} .section-title{{display:flex;align-items:center;gap:18px;margin-bottom:28px}} .section-title>span{{font-size:34px;color:#b9dce6;font-weight:800}} .section-title small{{color:var(--teal);font-weight:800;letter-spacing:.12em}} .section-title h2{{margin:2px 0;color:var(--navy);font-size:28px}}
    h3{{color:var(--navy);margin-top:28px}} .kpi-grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:16px 0 28px}} .kpi{{border:1px solid var(--line);border-radius:12px;padding:18px;background:#fbfdfe}} .kpi span,.kpi small{{display:block;color:var(--muted);font-size:11px}} .kpi b{{display:block;font-size:28px;color:var(--navy);margin:8px 0 2px}}
    .insight-grid{{display:grid;grid-template-columns:1fr 1fr;gap:12px}} .insight{{border:1px solid var(--line);border-left:4px solid var(--teal);border-radius:10px;padding:14px}} .insight.critical{{border-left-color:{BRAND['danger']}}}.insight.attention{{border-left-color:{BRAND['warning']}}}.insight.positive{{border-left-color:{BRAND['success']}}}.insight p{{margin:5px 0;color:var(--muted)}}
    .recs{{padding:0;list-style:none}} .recs li{{display:flex;gap:14px;border-bottom:1px solid #edf2f4;padding:13px 0}} .priority{{font-size:9px;font-weight:800;color:var(--teal);min-width:64px}} .recs small{{display:block;color:var(--muted);margin-top:3px}}
    .table-wrap{{overflow:auto;border:1px solid var(--line);border-radius:10px}} table{{border-collapse:collapse;width:100%;font-size:12px}} th{{background:#edf7fa;color:var(--navy);text-align:left;padding:10px;border-bottom:1px solid #cfe2e8;white-space:nowrap}} td{{padding:9px 10px;border-bottom:1px solid #edf2f4;vertical-align:top}} tr:nth-child(even) td{{background:#fbfcfd}}
    .score-ring{{width:142px;height:142px;border-radius:50%;border:14px solid #d8eef4;display:flex;flex-direction:column;align-items:center;justify-content:center;margin:0 0 24px}} .score-ring b{{font-size:34px;color:var(--navy)}} .score-ring span,.score-ring small{{font-size:10px;color:var(--muted)}}
    .figure-grid{{display:grid;gap:20px}} figure{{margin:0;border:1px solid var(--line);border-radius:12px;padding:16px;background:#fff}} figure svg{{width:100%;height:auto}} .figure-head{{display:flex;gap:12px;align-items:center;margin-bottom:10px}} .figure-head span{{font-size:10px;color:var(--teal);font-weight:800;text-transform:uppercase}} figcaption{{font-size:10px;color:var(--muted);margin-top:8px}}
    .question{{background:#edf7fa;border-left:4px solid var(--teal);padding:14px 16px;border-radius:8px}} .question span{{display:block;color:var(--teal);font-size:10px;font-weight:800;text-transform:uppercase}} .question b{{display:block;margin-top:4px}} .answer{{padding:18px 0;font-size:15px}}
    .finding{{display:flex;gap:12px;padding:13px 0;border-bottom:1px solid #edf2f4}} .finding>span{{font-size:9px;color:var(--teal);font-weight:800;min-width:70px}} .finding p{{margin:4px 0;color:var(--muted)}}
    .method-list{{display:grid;gap:10px}} .method-list>div{{display:grid;grid-template-columns:38px 1fr;gap:12px;align-items:start;border:1px solid var(--line);padding:12px;border-radius:9px}} .method-list span{{color:var(--teal);font-weight:800}} .method-list p{{margin:0}}
    .provenance{{display:grid;grid-template-columns:1fr 1fr;gap:1px;background:var(--line);border:1px solid var(--line)}} .provenance div{{background:white;padding:11px 13px}} .provenance span{{display:block;font-size:9px;text-transform:uppercase;color:var(--muted)}} .provenance b{{font-size:11px;overflow-wrap:anywhere}}
    .empty{{color:var(--muted)}} @media(max-width:760px){{.document{{margin:0}}.cover,.toc,section{{padding:34px 24px}}h1{{font-size:36px}}.cover-meta,.kpi-grid,.insight-grid,.toc-grid,.provenance{{grid-template-columns:1fr}}}}
    @media print{{body{{background:#fff}}.document{{margin:0;box-shadow:none;max-width:none}}.cover,.toc,section{{break-after:auto}}.cover{{break-after:page}}.toc{{break-after:page}}a{{color:inherit}}}}

    .lead{{font-size:16px;line-height:1.65;color:#526b77;max-width:900px}}.story-grid{{display:grid;grid-template-columns:1fr 1fr;gap:12px;margin-top:18px}}.story-card{{display:grid;grid-template-columns:38px 1fr;gap:12px;border:1px solid var(--line);border-left:4px solid var(--teal);border-radius:8px;padding:14px;background:#fbfdfe}}.story-card.critical{{border-left-color:#b44a53}}.story-card.attention{{border-left-color:#b77a14}}.story-card.positive{{border-left-color:#178a63}}.story-rank{{width:30px;height:30px;border-radius:7px;background:#e9f6fa;color:#1686a4;display:grid;place-items:center;font-weight:800;font-size:11px}}.story-card b{{color:var(--navy)}}.story-card p{{margin:5px 0 8px;color:#516a76}}.story-card small{{display:block;margin-top:4px;color:#71858f}}.story-conclusion{{margin-top:18px;padding:15px 17px;background:#eef8fb;border-left:4px solid var(--teal);border-radius:7px}}.story-conclusion span{{font-size:10px;letter-spacing:.12em;font-weight:800;color:var(--teal)}}.story-conclusion p{{margin:5px 0 0}}.figure-reason,.figure-insight{{font-size:12px;color:#5b707b;margin:8px 0}}.figure-insight{{background:#f2f8fa;border-left:3px solid var(--teal);padding:9px 10px}}.limit-list{{display:grid;gap:10px}}.limit-list article{{border:1px solid var(--line);border-radius:8px;padding:13px 15px;background:#fbfcfd}}.limit-list b{{color:var(--navy)}}.limit-list p{{margin:5px 0;color:#5e7480}}.limit-list small{{color:#728690}}@media(max-width:760px){{.story-grid{{grid-template-columns:1fr}}}}
    """
    return f"<!doctype html><html lang='fr'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>{html.escape(report['title'])}</title><style>{css}</style></head><body><main class='document'><section class='cover'><div class='brand'><div class='brand-mark'>DV</div><div>DataVision AI</div></div><div class='cover-main'><div class='cover-kicker'>{html.escape(report.get('template_label','Rapport analytique'))}</div><h1>{html.escape(report['title'])}</h1><div class='subtitle'>{html.escape(report.get('subtitle',''))}</div><div class='cover-meta'><div><span>Dataset</span><b>{html.escape(str(report['dataset']['name']))}</b></div><div><span>Version</span><b>v{html.escape(str(report['dataset']['version']))}</b></div><div><span>Date</span><b>{html.escape(_format_date(report.get('created_at')))}</b></div></div></div><div class='cover-foot'><span>{html.escape(report.get('organization') or report.get('author') or 'DataVision AI')}</span><span>Rapport reproductible</span></div></section><section class='toc'><h2>Sommaire</h2><div class='toc-grid'>{toc}</div></section>{''.join(sections)}</main></body></html>"


# ----------------------------- DOCX -----------------------------

def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.replace("#", ""))


def _set_cell_border(cell, color: str = "D8E4E9"):
    tc_pr = cell._tc.get_or_add_tcPr(); borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tc_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"; element = borders.find(qn(tag))
        if element is None: element = OxmlElement(tag); borders.append(element)
        element.set(qn("w:val"), "single"); element.set(qn("w:sz"), "3"); element.set(qn("w:color"), color)


def _docx_style(doc: Document):
    section = doc.sections[0]; section.top_margin = Inches(.65); section.bottom_margin = Inches(.65); section.left_margin = Inches(.72); section.right_margin = Inches(.72); section.different_first_page_header_footer = True; section.first_page_header.paragraphs[0].text = ""; section.first_page_footer.paragraphs[0].text = ""
    normal = doc.styles["Normal"]; normal.font.name = "Aptos"; normal.font.size = Pt(9.5); normal.font.color.rgb = RGBColor(36, 56, 68)
    normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.12
    for name, size, color, before, after in [("Title", 30, BRAND["navy"], 0, 10), ("Heading 1", 20, BRAND["navy"], 16, 8), ("Heading 2", 14, BRAND["teal_dark"], 12, 6), ("Heading 3", 11, BRAND["navy"], 10, 4)]:
        style = doc.styles[name]; style.font.name = "Aptos Display"; style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color.replace("#", "")); style.font.bold = True; style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
    hdr = section.header.paragraphs[0]; hdr.text = "DataVision AI    |    Data Intelligence Report"; hdr.style = doc.styles["Normal"]; hdr.runs[0].font.size = Pt(8); hdr.runs[0].font.color.rgb = RGBColor.from_string(BRAND["muted"].replace("#", ""))
    f = section.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER; r=f.add_run("DataVision AI - rapport reproductible    |    Page "); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(BRAND["muted"].replace("#", "")); fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); f._p.append(fld)


def _docx_kpis(doc: Document, kpis: list[dict[str, Any]]):
    if not kpis: return
    table=doc.add_table(rows=1, cols=len(kpis)); table.autofit=True
    for i,k in enumerate(kpis):
        cell=table.cell(0,i); _set_cell_shading(cell,"EFF8FB"); _set_cell_border(cell,"CFE4EA"); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        a=p.add_run(str(k.get("label",""))); a.font.size=Pt(8); a.font.color.rgb=RGBColor.from_string(BRAND["muted"].replace("#", "")); a.bold=True
        p.add_run("\n"); b=p.add_run(_fmt(k.get("value"))); b.font.size=Pt(18); b.font.color.rgb=RGBColor.from_string(BRAND["navy"].replace("#", "")); b.bold=True
        p.add_run("\n"); c=p.add_run(str(k.get("hint",""))); c.font.size=Pt(7); c.font.color.rgb=RGBColor.from_string(BRAND["muted"].replace("#", ""))
    doc.add_paragraph()


def _docx_table(doc: Document, columns: list[str], rows: list[dict[str, Any]], max_rows: int | None = None):
    if max_rows is not None: rows=rows[:max_rows]
    if not columns: return
    table=doc.add_table(rows=1, cols=len(columns)); table.style="Table Grid"
    for i,c in enumerate(columns):
        cell=table.rows[0].cells[i]; cell.text=str(c); _set_cell_shading(cell,"DDEFF4"); _set_cell_border(cell)
        for run in cell.paragraphs[0].runs: run.font.bold=True; run.font.size=Pt(7.5); run.font.color.rgb=RGBColor.from_string(BRAND["navy"].replace("#", ""))
    for idx,row in enumerate(rows):
        cells=table.add_row().cells
        for i,c in enumerate(columns):
            cells[i].text=_fmt(row.get(c)); _set_cell_border(cells[i]);
            if idx%2: _set_cell_shading(cells[i],"F8FBFC")
            for run in cells[i].paragraphs[0].runs: run.font.size=Pt(7.2)
    doc.add_paragraph()


def _docx_cover(doc: Document, report: dict[str, Any]):
    for _ in range(3): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("DV"); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=RGBColor.from_string(BRAND["teal"].replace("#", ""))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("DATAVISION AI"); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string(BRAND["navy"].replace("#", ""))
    doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(report["title"]); r.bold=True; r.font.size=Pt(30); r.font.color.rgb=RGBColor.from_string(BRAND["navy"].replace("#", ""))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(report.get("subtitle", "")); r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(BRAND["muted"].replace("#", ""))
    doc.add_paragraph(); table=doc.add_table(rows=1,cols=3)
    meta=[("Dataset",report["dataset"]["name"]),("Version",f"v{report['dataset']['version']}"),("Genere",_format_date(report.get("created_at")))]
    for i,(label,value) in enumerate(meta):
        cell=table.cell(0,i); _set_cell_shading(cell,"EFF8FB"); _set_cell_border(cell,"D5E6EB"); p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        a=p.add_run(label.upper()+"\n"); a.font.size=Pt(7); a.bold=True; a.font.color.rgb=RGBColor.from_string(BRAND["teal_dark"].replace("#", "")); b=p.add_run(str(value)); b.font.size=Pt(9); b.bold=True
    for _ in range(6): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(report.get("organization") or report.get("author") or "DataVision AI"); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(BRAND["muted"].replace("#", ""))
    doc.add_page_break()


def _docx_export(report: dict[str, Any], path: Path):
    doc=Document(); _docx_style(doc); _docx_cover(doc,report)
    doc.add_heading("Sommaire", level=1)
    for e in report.get("section_outline",[]):
        p=doc.add_paragraph(); p.style=doc.styles["Normal"]; a=p.add_run(f"{e['number']:02d}  "); a.bold=True; a.font.color.rgb=RGBColor.from_string(BRAND["teal"].replace("#", "")); b=p.add_run(e["title"]); b.bold=True
    doc.add_page_break()
    with tempfile.TemporaryDirectory(prefix="datavision-report-") as tmpdir:
        tmp=Path(tmpdir); sec=0
        for block in report.get("blocks",[]):
            sec+=1; doc.add_heading(f"{sec}. {block.get('title','Section')}",level=1); kind=block.get("type")
            if kind=="executive_summary":
                d=block["data"]; _docx_kpis(doc,d.get("kpis",[])); doc.add_heading("Points cles",level=2)
                for h in d.get("highlights",[]):
                    p=doc.add_paragraph(style="List Bullet"); a=p.add_run(h.get("title","")+" - "); a.bold=True; p.add_run(h.get("text",""))
                doc.add_heading("Recommandations",level=2)
                for r in d.get("recommendations",[]):
                    p=doc.add_paragraph(style="List Number"); a=p.add_run(str(r.get("priority","medium")).upper()+" - "); a.bold=True; a.font.color.rgb=RGBColor.from_string(BRAND["teal_dark"].replace("#", "")); p.add_run(r.get("action","")); p.add_run(f" ({r.get('evidence','')})")
            elif kind=="analytical_story":
                d=block["data"]; doc.add_paragraph(d.get("opening","")); doc.add_heading("Constats prioritaires",level=2)
                for f in d.get("findings",[]):
                    table=doc.add_table(rows=1,cols=2); table.autofit=False; table.columns[0].width=Inches(.5); table.columns[1].width=Inches(6.0); left=table.cell(0,0); right=table.cell(0,1); left.width=Inches(.5); right.width=Inches(6.0); _set_cell_shading(left,"EAF6F9"); _set_cell_border(left,"CFE4EA"); _set_cell_border(right,"DCE7EB")
                    lp=left.paragraphs[0]; lp.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=lp.add_run(f"{int(f.get('rank',0)):02d}"); rr.bold=True; rr.font.color.rgb=RGBColor.from_string(BRAND["teal_dark"].replace("#",""))
                    rp=right.paragraphs[0]; a=rp.add_run(str(f.get("title",""))+"\n"); a.bold=True; a.font.color.rgb=RGBColor.from_string(BRAND["navy"].replace("#","")); rp.add_run(str(f.get("statement",""))+"\n"); ev=rp.add_run("Preuve : "+str(f.get("evidence",""))+"\n"); ev.italic=True; ev.font.size=Pt(8); it=rp.add_run("Lecture : "+str(f.get("interpretation",""))); it.font.size=Pt(8); it.font.color.rgb=RGBColor.from_string(BRAND["muted"].replace("#","")); doc.add_paragraph()
                doc.add_heading("Conclusion de lecture",level=2); doc.add_paragraph(d.get("conclusion",""))
            elif kind=="overview":
                d=block["data"]; _docx_kpis(doc,[{"label":"Observations","value":d.get("rows"),"hint":"lignes"},{"label":"Variables","value":d.get("columns"),"hint":"colonnes"},{"label":"Doublons","value":d.get("duplicates"),"hint":"lignes"},{"label":"Qualite","value":d.get("quality_score"),"hint":"/100"}]);
                if d.get("type_summary"): doc.add_heading("Structure des variables",level=2); _docx_table(doc,["type","count"],d.get("type_summary",[]))
                if d.get("missing_summary"): doc.add_heading("Variables les plus incompletes",level=2); _docx_table(doc,["variable","missing","missing_pct"],d.get("missing_summary",[]))
            elif kind=="quality":
                d=block["data"]; p=doc.add_paragraph(); a=p.add_run(f"Score de qualite : {_fmt(d.get('score'))}/100"); a.bold=True; a.font.size=Pt(14); a.font.color.rgb=RGBColor.from_string(BRAND["teal_dark"].replace("#", "")); doc.add_paragraph(f"{d.get('issues_count',0)} probleme(s) detecte(s).")
                rows=[{"severity":i.get("severity"),"column":i.get("column","General"),"description":i.get("description"),"recommendation":i.get("recommendation")} for i in d.get("issues",[])]; _docx_table(doc,["severity","column","description","recommendation"],rows)
            elif kind=="table": _docx_table(doc,block.get("columns",[]),block.get("rows",[]))
            elif kind=="visualizations":
                items=block.get("items",[])
                if not items: doc.add_paragraph("Aucune visualisation epinglee pour cette version du dataset.")
                for idx,item in enumerate(items,1):
                    doc.add_heading(f"Figure {idx}. {item.get('title','Visualisation')}",level=2); img=tmp/f"chart-{idx}.png"
                    if _drawing_png(item,img):
                        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(img),width=Inches(6.55))
                    viz=item.get("visualization",{}); p=doc.add_paragraph(f"Type : {viz.get('type','-')} - dataset v{item.get('dataset_version','-')}"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs: run.font.size=Pt(7.5); run.font.color.rgb=RGBColor.from_string(BRAND["muted"].replace("#", ""))
                    if item.get("reason"):
                        p=doc.add_paragraph(); a=p.add_run("Objectif : "); a.bold=True; a.font.color.rgb=RGBColor.from_string(BRAND["teal_dark"].replace("#","")); p.add_run(str(item.get("reason")))
                    if item.get("insight"):
                        p=doc.add_paragraph(); a=p.add_run("Lecture : "); a.bold=True; a.font.color.rgb=RGBColor.from_string(BRAND["teal_dark"].replace("#","")); p.add_run(str(item.get("insight")))
            elif kind=="ai_analysis":
                d=block["data"]; p=doc.add_paragraph(); a=p.add_run("Question analytique\n"); a.bold=True; a.font.color.rgb=RGBColor.from_string(BRAND["teal_dark"].replace("#", "")); p.add_run(d.get("question","")); doc.add_paragraph(d.get("answer","")); doc.add_heading("Constats verifies",level=2)
                for finding in d.get("findings",[]):
                    p=doc.add_paragraph(style="List Bullet"); a=p.add_run(finding.get("title","")+" - "); a.bold=True; p.add_run(finding.get("statement",""))
            elif kind=="limitations":
                for item in block.get("items",[]):
                    p=doc.add_paragraph(); a=p.add_run(str(item.get("title",""))+" - "); a.bold=True; a.font.color.rgb=RGBColor.from_string(BRAND["navy"].replace("#","")); p.add_run(str(item.get("text",""))); p.add_run("\n"); b=p.add_run("Prudence : "+str(item.get("mitigation",""))); b.italic=True; b.font.size=Pt(8); b.font.color.rgb=RGBColor.from_string(BRAND["muted"].replace("#",""))
            elif kind=="methodology":
                for i,item in enumerate(block.get("items",[]),1):
                    p=doc.add_paragraph(); a=p.add_run(f"{i:02d}  "); a.bold=True; a.font.color.rgb=RGBColor.from_string(BRAND["teal"].replace("#", "")); p.add_run(item)
            elif kind=="provenance":
                rows=[{"element":k,"valeur":_fmt(v)} for k,v in block.get("data",{}).items()]; _docx_table(doc,["element","valeur"],rows)
            else: doc.add_paragraph(block.get("text",""))
            if sec < len(report.get("blocks",[])): doc.add_page_break()
    doc.save(path)


# ----------------------------- PDF -----------------------------

def _pdf_styles():
    base=getSampleStyleSheet()
    return {
        "title":ParagraphStyle("DVTitle",parent=base["Title"],fontName="Helvetica-Bold",fontSize=28,leading=31,textColor=_color(BRAND["navy"]),alignment=TA_LEFT,spaceAfter=10),
        "subtitle":ParagraphStyle("DVSubtitle",parent=base["BodyText"],fontName="Helvetica",fontSize=13,leading=18,textColor=_color(BRAND["muted"]),spaceAfter=14),
        "h1":ParagraphStyle("DVH1",parent=base["Heading1"],fontName="Helvetica-Bold",fontSize=19,leading=22,textColor=_color(BRAND["navy"]),spaceBefore=6,spaceAfter=12),
        "h2":ParagraphStyle("DVH2",parent=base["Heading2"],fontName="Helvetica-Bold",fontSize=12,leading=15,textColor=_color(BRAND["teal_dark"]),spaceBefore=10,spaceAfter=6),
        "body":ParagraphStyle("DVBody",parent=base["BodyText"],fontName="Helvetica",fontSize=8.5,leading=12,textColor=_color(BRAND["ink"]),spaceAfter=5),
        "small":ParagraphStyle("DVSmall",parent=base["BodyText"],fontName="Helvetica",fontSize=7,leading=9,textColor=_color(BRAND["muted"]),spaceAfter=3),
        "center":ParagraphStyle("DVCenter",parent=base["BodyText"],fontName="Helvetica",fontSize=8,leading=11,textColor=_color(BRAND["muted"]),alignment=TA_CENTER),
    }


def _pdf_header_footer(canvas, doc, report: dict[str, Any]):
    canvas.saveState(); w,h=A4
    canvas.setStrokeColor(_color("#DDE7EB")); canvas.setLineWidth(.5); canvas.line(20*mm,h-15*mm,w-20*mm,h-15*mm)
    canvas.setFillColor(_color(BRAND["muted"])); canvas.setFont("Helvetica",7); canvas.drawString(20*mm,h-11*mm,"DataVision AI | Data Intelligence Report"); canvas.drawRightString(w-20*mm,h-11*mm,f"{report['dataset']['name']} - v{report['dataset']['version']}")
    canvas.line(20*mm,13*mm,w-20*mm,13*mm); canvas.drawString(20*mm,8.5*mm,"Rapport reproductible"); canvas.drawRightString(w-20*mm,8.5*mm,f"Page {doc.page}"); canvas.restoreState()


def _pdf_cover(canvas, doc, report: dict[str, Any]):
    canvas.saveState(); w,h=A4
    canvas.setFillColor(_color("#F3FAFC")); canvas.rect(0,0,w,h,fill=1,stroke=0)
    canvas.setFillColor(_color(BRAND["teal"])); canvas.rect(0,0,15*mm,h,fill=1,stroke=0)
    canvas.roundRect(28*mm,h-43*mm,18*mm,18*mm,3*mm,fill=0,stroke=1); canvas.setFillColor(_color(BRAND["teal"])); canvas.setFont("Helvetica-Bold",16); canvas.drawCentredString(37*mm,h-36*mm,"DV")
    canvas.setFillColor(_color(BRAND["navy"])); canvas.setFont("Helvetica-Bold",10); canvas.drawString(51*mm,h-33*mm,"DATAVISION AI")
    canvas.setFillColor(_color(BRAND["teal_dark"])); canvas.setFont("Helvetica-Bold",8); canvas.drawString(28*mm,h-70*mm,str(report.get("template_label","Rapport analytique")).upper())
    title=report["title"]; canvas.setFillColor(_color(BRAND["navy"])); canvas.setFont("Helvetica-Bold",27)
    words=title.split(); lines=[]; cur=""
    for word in words:
        test=(cur+" "+word).strip()
        if canvas.stringWidth(test,"Helvetica-Bold",27)>145*mm and cur: lines.append(cur); cur=word
        else: cur=test
    if cur: lines.append(cur)
    y=h-87*mm
    for line in lines[:3]: canvas.drawString(28*mm,y,line); y-=11*mm
    canvas.setFillColor(_color(BRAND["muted"])); canvas.setFont("Helvetica",12); canvas.drawString(28*mm,y-2*mm,report.get("subtitle","")[:105])
    y_meta=58*mm; labels=[("DATASET",report["dataset"]["name"]),("VERSION",f"v{report['dataset']['version']}"),("GENERE",_format_date(report.get("created_at")))]
    x=28*mm
    for lab,val in labels:
        canvas.setFillColor(_color("#E5F3F7")); canvas.roundRect(x,y_meta,48*mm,22*mm,2*mm,fill=1,stroke=0); canvas.setFillColor(_color(BRAND["teal_dark"])); canvas.setFont("Helvetica-Bold",6.5); canvas.drawString(x+4*mm,y_meta+14*mm,lab); canvas.setFillColor(_color(BRAND["ink"])); canvas.setFont("Helvetica-Bold",8); canvas.drawString(x+4*mm,y_meta+7*mm,str(val)[:25]); x+=52*mm
    canvas.setFillColor(_color(BRAND["muted"])); canvas.setFont("Helvetica",7.5); canvas.drawString(28*mm,24*mm,report.get("organization") or report.get("author") or "DataVision AI"); canvas.drawRightString(w-22*mm,24*mm,"Analyse verifiable - provenance conservee"); canvas.restoreState()


def _pdf_table(columns: list[str], rows: list[dict[str, Any]], col_widths=None, font_size=6.5):
    data=[[Paragraph(f"<b>{html.escape(str(c))}</b>",ParagraphStyle("th",fontName="Helvetica-Bold",fontSize=font_size,leading=font_size+2,textColor=_color(BRAND["navy"]))) for c in columns]]
    style=ParagraphStyle("td",fontName="Helvetica",fontSize=font_size,leading=font_size+2,textColor=_color(BRAND["ink"]));
    for row in rows: data.append([Paragraph(html.escape(_fmt(row.get(c))),style) for c in columns])
    t=Table(data,colWidths=col_widths,repeatRows=1,hAlign="LEFT"); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),_color("#DDEFF4")),("GRID",(0,0),(-1,-1),.3,_color("#D6E2E7")),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,_color("#FAFCFD")]),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)])); return t


def _pdf_kpis(kpis: list[dict[str, Any]], styles):
    cells=[]
    for k in kpis:
        cells.append([Paragraph(html.escape(k.get("label","")),styles["small"]),Paragraph(f"<b>{html.escape(_fmt(k.get('value')))}</b>",ParagraphStyle("kv",fontName="Helvetica-Bold",fontSize=17,leading=19,textColor=_color(BRAND["navy"]),alignment=TA_CENTER)),Paragraph(html.escape(k.get("hint","")),styles["small"])])
    nested=[Table([[p] for p in x],colWidths=[37*mm],style=[("BACKGROUND",(0,0),(-1,-1),_color("#F4FAFC")),("BOX",(0,0),(-1,-1),.4,_color("#D5E7EC")),("ALIGN",(0,0),(-1,-1),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]) for x in cells]
    return Table([nested],colWidths=[39*mm]*len(nested),hAlign="LEFT",style=[("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),3)])


def _pdf_export(report: dict[str, Any], path: Path):
    styles=_pdf_styles(); story=[]
    # Cover page is drawn by onFirstPage; invisible flowables reserve it then force page break.
    story.extend([Spacer(1,240*mm),PageBreak()])
    story.append(Paragraph("Sommaire",styles["h1"])); story.append(Spacer(1,4))
    toc=[]
    for e in report.get("section_outline",[]):
        toc.append([Paragraph(f"<b>{e['number']:02d}</b>",styles["body"]),Paragraph(html.escape(e["title"]),styles["body"])])
    if toc:
        t=Table(toc,colWidths=[12*mm,150*mm],hAlign="LEFT"); t.setStyle(TableStyle([("LINEBELOW",(0,0),(-1,-1),.25,_color("#E8EEF1")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TEXTCOLOR",(0,0),(0,-1),_color(BRAND["teal"])),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)])); story.append(t)
    story.append(PageBreak())
    sec=0
    for block in report.get("blocks",[]):
        sec+=1; story.append(Paragraph(f"{sec:02d}  {html.escape(block.get('title','Section'))}",styles["h1"])); kind=block.get("type")
        if kind=="executive_summary":
            d=block["data"]; story.append(_pdf_kpis(d.get("kpis",[]),styles)); story.append(Spacer(1,8)); story.append(Paragraph("Points cles",styles["h2"]))
            for h in d.get("highlights",[]):
                color=BRAND["danger"] if h.get("level")=="critical" else BRAND["warning"] if h.get("level")=="attention" else BRAND["success"] if h.get("level")=="positive" else BRAND["teal"]
                cell=Table([[Paragraph(f"<b>{html.escape(h.get('title',''))}</b><br/>{html.escape(h.get('text',''))}",styles["body"])]],colWidths=[160*mm],style=[("BACKGROUND",(0,0),(-1,-1),_color("#FAFCFD")),("LINEBEFORE",(0,0),(0,-1),3,_color(color)),("BOX",(0,0),(-1,-1),.3,_color("#DFE8EC")),("LEFTPADDING",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]); story.extend([cell,Spacer(1,4)])
            story.append(Paragraph("Recommandations",styles["h2"]));
            for i,r in enumerate(d.get("recommendations",[]),1): story.append(Paragraph(f"<b>{i}. {html.escape(str(r.get('priority','medium')).upper())}</b> - {html.escape(r.get('action',''))}<br/><font color='{BRAND['muted']}'>{html.escape(r.get('evidence',''))}</font>",styles["body"]))
        elif kind=="analytical_story":
            d=block["data"]; story.append(Paragraph(html.escape(d.get("opening","")),styles["body"])); story.append(Paragraph("Constats prioritaires",styles["h2"]))
            for f in d.get("findings",[]):
                sev=str(f.get("severity","info")); color=BRAND["danger"] if sev=="critical" else BRAND["warning"] if sev=="attention" else BRAND["success"] if sev=="positive" else BRAND["teal"]
                body=Paragraph(f"<b>{html.escape(str(f.get('title','')))}</b><br/>{html.escape(str(f.get('statement','')))}<br/><font size='7' color='{BRAND['muted']}'><b>Preuve :</b> {html.escape(str(f.get('evidence','')))}<br/><b>Lecture :</b> {html.escape(str(f.get('interpretation','')))}</font>",styles["body"]); rank=Paragraph(f"<b>{int(f.get('rank',0)):02d}</b>",styles["center"]); card=Table([[rank,body]],colWidths=[12*mm,148*mm],style=[("BACKGROUND",(0,0),(-1,-1),_color("#FAFCFD")),("LINEBEFORE",(0,0),(0,-1),3,_color(color)),("BOX",(0,0),(-1,-1),.3,_color("#DFE8EC")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]); story.extend([card,Spacer(1,4)])
            story.extend([Paragraph("Conclusion de lecture",styles["h2"]),Paragraph(html.escape(d.get("conclusion","")),styles["body"])])
        elif kind=="overview":
            d=block["data"]; story.append(_pdf_kpis([{"label":"Observations","value":d.get("rows"),"hint":"lignes"},{"label":"Variables","value":d.get("columns"),"hint":"colonnes"},{"label":"Doublons","value":d.get("duplicates"),"hint":"lignes"},{"label":"Qualite","value":d.get("quality_score"),"hint":"/100"}],styles)); story.append(Spacer(1,8))
            if d.get("type_summary"): story.extend([Paragraph("Structure des variables",styles["h2"]),_pdf_table(["type","count"],d["type_summary"],[75*mm,35*mm],7),Spacer(1,8)])
            if d.get("missing_summary"): story.extend([Paragraph("Variables les plus incompletes",styles["h2"]),_pdf_table(["variable","missing","missing_pct"],d["missing_summary"],[70*mm,35*mm,35*mm],7)])
        elif kind=="quality":
            d=block["data"]; story.append(Paragraph(f"<b>Score de qualite : {_fmt(d.get('score'))}/100</b> - {d.get('issues_count',0)} probleme(s) detecte(s).",styles["body"])); rows=[{"severity":i.get("severity"),"column":i.get("column","General"),"description":i.get("description"),"recommendation":i.get("recommendation")} for i in d.get("issues",[])]; story.append(_pdf_table(["severity","column","description","recommendation"],rows,[18*mm,28*mm,55*mm,60*mm],5.8))
        elif kind=="table": story.append(_pdf_table(block.get("columns",[]),block.get("rows",[]),font_size=5.4))
        elif kind=="visualizations":
            items=block.get("items",[])
            if not items: story.append(Paragraph("Aucune visualisation epinglee pour cette version du dataset.",styles["body"]))
            for idx,item in enumerate(items,1):
                figure=[Paragraph(f"Figure {idx}. {html.escape(item.get('title','Visualisation'))}",styles["h2"])]
                drawing=_chart_drawing(item,450,225)
                if drawing is not None: figure.append(drawing)
                viz=item.get("visualization",{}); figure.append(Paragraph(f"Type : {html.escape(str(viz.get('type','-')))} - dataset v{html.escape(str(item.get('dataset_version','-')))}",styles["center"]))
                if item.get("reason"): figure.append(Paragraph(f"<b>Objectif :</b> {html.escape(str(item.get('reason')))}",styles["body"]))
                if item.get("insight"): figure.append(Paragraph(f"<b>Lecture :</b> {html.escape(str(item.get('insight')))}",styles["body"]))
                figure.append(Spacer(1,8)); story.append(KeepTogether(figure))
        elif kind=="ai_analysis":
            d=block["data"]; q=Table([[Paragraph("QUESTION ANALYTIQUE",styles["small"]),Paragraph(f"<b>{html.escape(d.get('question',''))}</b>",styles["body"])]],colWidths=[35*mm,125*mm],style=[("BACKGROUND",(0,0),(-1,-1),_color("#EDF7FA")),("LINEBEFORE",(0,0),(0,-1),3,_color(BRAND["teal"])),("BOX",(0,0),(-1,-1),.3,_color("#D7E6EB")),("VALIGN",(0,0),(-1,-1),"TOP"),]); story.extend([q,Spacer(1,8),Paragraph(html.escape(d.get("answer","")),styles["body"]),Paragraph("Constats verifies",styles["h2"])]);
            for f in d.get("findings",[]): story.append(Paragraph(f"<b>{html.escape(f.get('title',''))}</b> - {html.escape(f.get('statement',''))}",styles["body"]))
        elif kind=="limitations":
            for item in block.get("items",[]):
                story.append(Paragraph(f"<b>{html.escape(str(item.get('title','')))}</b> - {html.escape(str(item.get('text','')))}<br/><font size='7' color='{BRAND['muted']}'><b>Prudence :</b> {html.escape(str(item.get('mitigation','')))}</font>",styles["body"]))
        elif kind=="methodology":
            for i,item in enumerate(block.get("items",[]),1): story.append(Paragraph(f"<b>{i:02d}</b>  {html.escape(item)}",styles["body"]))
        elif kind=="provenance": story.append(_pdf_table(["element","valeur"],[{"element":k,"valeur":_fmt(v)} for k,v in block.get("data",{}).items()],[45*mm,115*mm],6.3))
        else: story.append(Paragraph(html.escape(block.get("text","")),styles["body"]))
        if sec<len(report.get("blocks",[])): story.append(PageBreak())
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=20*mm,leftMargin=20*mm,topMargin=20*mm,bottomMargin=18*mm,title=report["title"],author=report.get("author") or "DataVision AI")
    doc.build(story,onFirstPage=lambda c,d:_pdf_cover(c,d,report),onLaterPages=lambda c,d:_pdf_header_footer(c,d,report))


def export_report(report_id: str, fmt: str) -> Path:
    report=get_report(report_id); fmt=fmt.lower(); base=_dir()/f"{_slug(report['title'])}-{report_id[:8]}"
    if fmt in {"md","markdown"}: path=base.with_suffix(".md"); path.write_text(_markdown(report),encoding="utf-8"); return path
    if fmt=="html": path=base.with_suffix(".html"); path.write_text(_html(report),encoding="utf-8"); return path
    if fmt=="docx": path=base.with_suffix(".docx"); _docx_export(report,path); return path
    if fmt=="pdf": path=base.with_suffix(".pdf"); _pdf_export(report,path); return path
    raise ValueError("Format d'export non supporte")
